"""把「酒吧数据模板.xlsx」转换成网页用的数据文件。

用法：python3 tools/build_data.py 酒吧数据模板.xlsx
输出：data/cities.json（城市列表）和 data/<城市拼音>.json（每个城市一个文件）

只导出页面要显示的内容：核实日期不导出；禁烟只导出“是否全店禁烟”。
"""
import hashlib, json, re, sys
from pathlib import Path
from openpyxl import load_workbook

# 城市中文名 → 文件名。新增城市时在这里加一行。
CITY_SLUGS = {"新乡": "xinxiang"}

TYPES = {"清吧", "啤酒吧"}
SMOKE = {"全店禁烟", "有吸烟区", "可吸烟", "未确认", ""}
NOISE = {"", "有驻唱", "音乐声过大"}
LINK_HOSTS = {
    "小红书链接": ("xiaohongshu.com", "xhslink.com"),
    "抖音链接": ("douyin.com", "iesdouyin.com"),
    "找门帖子链接": ("xiaohongshu.com", "xhslink.com", "douyin.com", "iesdouyin.com"),
}
UNKNOWN_HOURS = re.compile(r"未明确|未知|待确认|不确定")

def cell(v):
    return "" if v is None else str(v).strip()

def check_link(col, url, errs, where):
    if not url:
        return ""
    m = re.match(r"^https://([^/?#]+)", url)
    host = m.group(1).lower() if m else ""
    if not host or not any(host == h or host.endswith("." + h) for h in LINK_HOSTS[col]):
        errs.append(f"{where}「{col}」不是有效的 https 链接或网站不对：{url}")
        return ""
    return url

def main(path):
    ws = load_workbook(path, read_only=True)["酒吧数据"]
    rows = list(ws.iter_rows(values_only=True))
    head = [cell(h) for h in rows[0]]
    col = {h.split("（")[0]: i for i, h in enumerate(head)}
    errs, cities, seen = [], {}, set()
    for n, r in enumerate(rows[1:], start=2):
        g = lambda k: cell(r[col[k]]) if col.get(k) is not None and col[k] < len(r) else ""
        name = g("店名")
        if not name:
            continue
        where = f"第 {n} 行（{name}）"
        city, typ, addr = g("城市"), g("类型"), g("地址")
        if city not in CITY_SLUGS:
            errs.append(f"{where}城市「{city}」还没有文件名，请在 tools/build_data.py 的 CITY_SLUGS 里加上"); continue
        if typ not in TYPES: errs.append(f"{where}类型只能是 清吧 / 啤酒吧，现在是「{typ}」")
        if not addr: errs.append(f"{where}缺少地址")
        smoke, noise = g("禁烟"), g("吵闹提醒")
        if smoke not in SMOKE: errs.append(f"{where}禁烟一栏的值「{smoke}」不在下拉选项里")
        if noise not in NOISE: errs.append(f"{where}吵闹提醒的值「{noise}」不在下拉选项里")
        bar = {
            "id": hashlib.sha1(f"{city}|{name}|{addr}".encode()).hexdigest()[:8],
            "name": name, "type": typ, "district": g("区"), "addr": addr,
            "hours": "" if UNKNOWN_HOURS.search(g("营业时间")) else g("营业时间"),
        }
        ll = g("坐标").replace("，", ",").replace(" ", "")
        if ll:
            try:
                lng, lat = (float(x) for x in ll.split(","))
                assert 73 < lng < 136 and 3 < lat < 54
                bar["lnglat"] = [lng, lat]
            except Exception:
                errs.append(f"{where}坐标格式不对（要求“经度,纬度”，在中国范围内）：{ll}")
        if smoke == "全店禁烟": bar["smokefree"] = True
        if noise: bar["noise"] = noise
        for c, key in (("小红书链接", "xhs"), ("抖音链接", "douyin"), ("找门帖子链接", "guide")):
            u = check_link(c, g(c), errs, where)
            if u: bar[key] = u
        if bar["id"] in seen: errs.append(f"{where}和前面某一行重复（城市、店名、地址都相同）")
        seen.add(bar["id"])
        cities.setdefault(city, []).append(bar)
    if errs:
        print("有 %d 处需要修改，数据文件没有生成：" % len(errs))
        print("\n".join("  - " + e for e in errs)); sys.exit(1)
    out = Path(__file__).resolve().parent.parent / "data"
    out.mkdir(exist_ok=True)
    index = []
    for city, bars in cities.items():
        slug = CITY_SLUGS[city]
        (out / f"{slug}.json").write_text(json.dumps(bars, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        index.append({"slug": slug, "name": city, "count": len(bars)})
    (out / "cities.json").write_text(json.dumps(index, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    for c in index: print(f"{c['name']}：{c['count']} 家 → data/{c['slug']}.json")

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "酒吧数据模板.xlsx")
